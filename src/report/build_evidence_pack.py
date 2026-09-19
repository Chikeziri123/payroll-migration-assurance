"""
Reconciliation evidence pack.

The artefact a migration programme hands to a steering group before
go-live. Every figure in it comes from a query against the recon schema,
so the pack regenerates on every pipeline run rather than being assembled
by hand and going stale.

Structure is deliberate: a sign-off summary that a programme board reads,
then the full detail behind it that a migration lead works from and an
auditor can check the summary against. A summary nobody can verify is a
claim, not evidence.

Writes to output/reconciliation-evidence-pack.xlsx
"""

import os
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from dotenv import load_dotenv
from sqlalchemy import create_engine

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "output"
OUT_FILE = OUT_DIR / "reconciliation-evidence-pack.xlsx"


def connect():
    """Build an engine from .env, accepting the usual variable spellings."""
    load_dotenv(ROOT / ".env")

    def pick(*names, default=None):
        for n in names:
            v = os.getenv(n)
            if v:
                return v
        return default

    host = pick("PGHOST", "DB_HOST", "POSTGRES_HOST", default="localhost")
    port = pick("PGPORT", "DB_PORT", "POSTGRES_PORT", default="5432")
    db = pick("PGDATABASE", "DB_NAME", "POSTGRES_DB", default="payrollmig")
    user = pick("PGUSER", "DB_USER", "POSTGRES_USER")
    pwd = pick("PGPASSWORD", "DB_PASSWORD", "POSTGRES_PASSWORD")

    if not user or not pwd:
        sys.exit(
            "Could not find database credentials in .env.\n"
            "Expected a user and password under one of: PGUSER/PGPASSWORD, "
            "DB_USER/DB_PASSWORD, POSTGRES_USER/POSTGRES_PASSWORD."
        )

    return create_engine(f"postgresql+psycopg2://{user}:{pwd}@{host}:{port}/{db}")


# Detail sheets. Order matters: the reader moves from programme level
# down to individual records, not the other way round.
DETAIL = [
    ("Record counts", "select * from recon.recon_record_counts order by infotype"),
    ("Exceptions", "select * from recon.recon_exceptions order by triage_rank, pernr"),
    ("Rejects", "select * from recon.int_rejects_resolved order by infotype, pernr"),
    ("Pay control", "select * from recon.recon_pay_current order by pernr"),
    ("Detection scores", "select * from profiling.detection_scores order by defect_code"),
    ("Table scores", "select * from profiling.table_scores order by weighted_score"),
    ("Completeness matrix", "select * from recon.recon_completeness_matrix order by pernr, infotype"),
]


def strip_timezones(frame):
    """
    Excel has no offset-aware datetime type, so timestamptz columns have
    to lose their offset before they can be written. Wall-clock time is
    kept rather than converted to UTC, because the pack is read by people
    in one place at one moment, not consumed by another system.
    """
    for col in frame.columns:
        if isinstance(frame[col].dtype, pd.DatetimeTZDtype):
            frame[col] = frame[col].dt.tz_localize(None)
    return frame


def build_summary(frames):
    """
    The sign-off sheet.

    Each line is a claim the programme is being asked to accept, with the
    figure that supports it and where in the pack that figure is proved.
    Nothing here is typed by hand.
    """
    counts = frames["Record counts"]
    exceptions = frames["Exceptions"]
    detection = frames["Detection scores"]

    infotypes_total = len(counts)
    infotypes_pass = int((counts["status"] == "PASS").sum())
    unexplained = int(counts["unexplained_absent"].sum())

    exc_total = len(exceptions)
    exc_unexplained = int((exceptions["recon_state"] == "UNEXPLAINED_ABSENT").sum())

    seeded = int(detection["seeded"].sum())
    detected = int(detection["detected"].sum())
    rate = (detected / seeded * 100) if seeded else 0.0

    rows = [
        ("Reconciliation", "Infotypes reconciling",
         f"{infotypes_pass} of {infotypes_total}",
         "PASS" if infotypes_pass == infotypes_total else "REVIEW",
         "Record counts"),
        ("Reconciliation", "Records lost without explanation",
         unexplained, "PASS" if unexplained == 0 else "FAIL",
         "Record counts"),
        ("Exceptions", "Exceptions raised", exc_total, "INFO", "Exceptions"),
        ("Exceptions", "Exceptions still unexplained",
         exc_unexplained, "PASS" if exc_unexplained == 0 else "FAIL",
         "Exceptions"),
        ("Data quality", "Seeded defects detected",
         f"{detected:,} of {seeded:,} ({rate:.1f}%)", "INFO", "Detection scores"),
    ]

    return pd.DataFrame(
        rows, columns=["Area", "Control", "Result", "Status", "Evidence sheet"]
    )


def autosize(worksheet, frame):
    """Width to content, capped, so nothing arrives as a column of hashes."""
    for i, col in enumerate(frame.columns, start=1):
        longest = max(
            [len(str(col))] + [len(str(v)) for v in frame[col].head(400).tolist()]
        )
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=i).column_letter
        ].width = min(max(longest + 2, 10), 55)


def ensure_writable(path):
    """
    Excel holds an exclusive lock on an open workbook, so a rebuild while
    the previous pack is open dies at the write step, after every query
    has already run. Check first and say so plainly.
    """
    if path.exists():
        try:
            with open(path, "a+b"):
                pass
        except PermissionError:
            sys.exit(
                f"\n{path.name} is open in Excel and cannot be overwritten.\n"
                "Close the workbook and run this again."
            )


def main():
    OUT_DIR.mkdir(exist_ok=True)
    ensure_writable(OUT_FILE)
    engine = connect()

    frames = {}
    for name, sql in DETAIL:
        print(f"  reading {name} ...", end=" ")
        frames[name] = strip_timezones(pd.read_sql(sql, engine))
        print(f"{len(frames[name]):,} rows")

    summary = build_summary(frames)

    with pd.ExcelWriter(OUT_FILE, engine="openpyxl") as writer:
        # The summary table goes in first so the sheet exists, then the
        # title block is written by hand above it. Going through a
        # DataFrame would turn the title into a column header and leave
        # the prose unmerged, which reads as a spreadsheet someone forgot
        # to finish rather than a document a board is asked to sign.
        summary.to_excel(writer, sheet_name="Sign-off summary",
                         index=False, startrow=7)

        for name in frames:
            frames[name].to_excel(writer, sheet_name=name[:31], index=False)

        for name, frame in frames.items():
            sheet = writer.sheets[name[:31]]
            sheet.freeze_panes = "A2"
            autosize(sheet, frame)

        ws = writer.sheets["Sign-off summary"]

        title_lines = [
            ("UK Payroll Migration Assurance Framework", True),
            ("Reconciliation evidence pack", False),
            (f"Generated {datetime.now():%d %B %Y at %H:%M}", False),
            ("", False),
            ("Every figure below is queried from the reconciliation layer "
             "at the moment of generation. The detail sheets behind this "
             "one hold the records each figure is derived from.", False),
        ]
        for i, (text, is_title) in enumerate(title_lines, start=1):
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
            cell = ws.cell(row=i, column=1, value=text)
            cell.font = Font(bold=is_title, size=14 if is_title else 11)
            cell.alignment = Alignment(vertical="center", wrap_text=not is_title)
        ws.row_dimensions[5].height = 30

        for cell in ws[8]:
            cell.font = Font(bold=True)

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 26
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 22

    print(f"\nWritten to {OUT_FILE}")


if __name__ == "__main__":
    main()